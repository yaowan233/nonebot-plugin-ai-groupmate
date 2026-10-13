"""Document extraction and deterministic table summaries; never execute file content."""

import io
import csv
import json
import math
import codecs
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from dataclasses import dataclass
from urllib.parse import urljoin

from bs4 import BeautifulSoup

from .file_workspace import FileTaskError

MAX_TEXT_CHARS = 1_000_000
MAX_TABLE_ROWS = 10_000
MAX_TABLE_COLUMNS = 100
MAX_UNPACKED_BYTES = 50 * 1024 * 1024


def decode_text(data: bytes, content_type: str = "") -> str:
    encodings = ["utf-16"] if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)) else ["utf-8-sig", "gb18030"]
    if "charset=" in content_type.lower():
        encodings.insert(0, content_type.lower().split("charset=", 1)[1].split(";", 1)[0].strip(' "'))
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except (UnicodeError, LookupError):
            continue
    raise FileTaskError("unsupported_encoding", "无法识别文本编码，请提供 UTF-8 文本。")


def bounded_text(text: str) -> str:
    if len(text) > MAX_TEXT_CHARS:
        raise FileTaskError("document_too_large", "文档提取后超过 100 万字符，请提供较小的文档。")
    return text


def extract_html(data: bytes, url: str, content_type: str = "") -> tuple[str, str, list[dict]]:
    soup = BeautifulSoup(decode_text(data, content_type), "html.parser")
    title = soup.title.get_text(" ", strip=True)[:300] if soup.title else ""
    for element in soup(["script", "style", "noscript", "template", "svg"]):
        element.decompose()
    links = []
    seen = set()
    for element in soup.find_all("a", href=True):
        link = urljoin(url, str(element["href"]))
        if link.startswith(("http://", "https://")) and len(link) <= 2048 and link not in seen:
            seen.add(link)
            links.append({"text": element.get_text(" ", strip=True)[:120], "url": link})
            if len(links) >= 500:
                break
    for row in soup.find_all("tr"):
        cells = row.find_all(["td", "th"], recursive=False)
        if cells:
            row.replace_with("\n" + " | ".join(cell.get_text(" ", strip=True) for cell in cells) + "\n")
    text = soup.get_text("\n", strip=True)
    if links:
        text += "\n\n[网页链接，最多 500 项]\n" + "\n".join(f"{link['text']}: {link['url']}" for link in links)
    return title, bounded_text(text), links


def validate_office_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_UNPACKED_BYTES:
                raise FileTaskError("document_too_large", "Office 文件解压后过大，无法读取。")
            if any(item.flag_bits & 1 for item in entries):
                raise FileTaskError("encrypted_document", "不支持加密的 Office 文件。")
            # defusedxml is installed for openpyxl; do not accept DTD/entity declarations in either format.
            for entry in entries:
                if entry.filename.endswith((".xml", ".rels")):
                    xml = archive.read(entry).upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise FileTaskError("unsafe_document", "文档包含不支持的 XML 实体声明。")
    except zipfile.BadZipFile:
        raise FileTaskError("invalid_document", "文件不是有效的 Office 文档。") from None


@dataclass
class Table:
    columns: list[str]
    rows: list[list]
    sheets: list[str]
    sheet: str


def read_table(data: bytes, name: str, sheet: str = "") -> Table:
    suffix = Path(name).suffix.lower()
    sheets: list[str] = []
    selected_sheet = ""
    workbook = None
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        validate_office_zip(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheets = workbook.sheetnames
        selected_sheet = sheet or sheets[0]
        if selected_sheet not in sheets:
            workbook.close()
            raise FileTaskError("sheet_not_found", "工作表不存在，可用工作表：" + ", ".join(sheets)[:500])
        worksheet = workbook[selected_sheet]
        if (worksheet.max_column or 0) > MAX_TABLE_COLUMNS or (worksheet.max_row or 0) > MAX_TABLE_ROWS + 1:
            workbook.close()
            raise FileTaskError("table_too_large", "工作表超过 100 列或 10000 行限制。")
        # Some producers incorrectly report A1:A1; let iteration discover real dimensions.
        worksheet.reset_dimensions()
        iterator = worksheet.iter_rows(values_only=True)
    elif suffix in {".csv", ".tsv"}:
        text = decode_text(data)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        iterator = csv.reader(io.StringIO(text, newline=""), dialect=dialect)
    else:
        raise FileTaskError("unsupported_table", "表格统计支持 CSV、TSV 和 XLSX。旧版 XLS 请先转换为 XLSX。")
    try:
        raw_rows = []
        char_count = 0
        for row in iterator:
            row = list(row)
            if len(row) > MAX_TABLE_COLUMNS or len(raw_rows) >= MAX_TABLE_ROWS + 1:
                raise FileTaskError("table_too_large", "表格超过 100 列或 10000 行限制。")
            char_count += sum(len(str(value)) for value in row if value is not None)
            if char_count > MAX_TEXT_CHARS:
                raise FileTaskError("table_too_large", "表格文字内容过大，请提供较小的表格。")
            raw_rows.append(row)
    finally:
        if workbook is not None:
            workbook.close()
    if not raw_rows:
        raise FileTaskError("empty_table", "表格为空。")
    width = max(map(len, raw_rows))
    headers = raw_rows[0] + [None] * (width - len(raw_rows[0]))
    columns = [str(value).strip() if value is not None and str(value).strip() else f"column_{index + 1}" for index, value in enumerate(headers)]
    if len(set(columns)) != len(columns):
        raise FileTaskError("duplicate_columns", "表格第一行包含重复列名，请提供列名唯一的表格。")
    rows = [row + [None] * (width - len(row)) for row in raw_rows[1:] if any(value not in (None, "") for value in row)]
    return Table(columns, rows, sheets, selected_sheet)


def table_text(table: Table) -> str:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(table.columns)
    writer.writerows(table.rows)
    return bounded_text(buffer.getvalue())


def extract_document(data: bytes, name: str, *, content_type: str = "", source_url: str = "", sheet: str = "") -> tuple[str, dict]:
    suffix = Path(name).suffix.lower()
    if suffix in {".html", ".htm"} or content_type.startswith("text/html"):
        title, text, _ = extract_html(data, source_url, content_type)
        return text, {"title": title}
    if suffix in {".csv", ".tsv", ".xlsx"}:
        table = read_table(data, name, sheet)
        return table_text(table), {"columns": table.columns, "row_count": len(table.rows), "sheets": table.sheets, "sheet": table.sheet,
                                   "formula_notice": "XLSX 公式保留为文本，不执行公式；统计工具不把公式当成计算后的数值。"}
    if suffix == ".pdf" or data.startswith(b"%PDF-"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise FileTaskError("encrypted_document", "不支持加密的 PDF。")
        if len(reader.pages) > 200:
            raise FileTaskError("document_too_large", "PDF 超过 200 页，请先拆分。")
        parts = []
        size = 0
        for index, page in enumerate(reader.pages):
            content = page.extract_text() or ""
            size += len(content)
            if size > MAX_TEXT_CHARS:
                raise FileTaskError("document_too_large", "PDF 提取文本过多，请先拆分。")
            parts.append(f"[第 {index + 1} 页]\n{content}")
        if size == 0:
            raise FileTaskError("ocr_required", "该 PDF 没有可提取的文本，可能是扫描件；当前工具不提供 OCR。")
        return bounded_text("\n\n".join(parts)), {"pages": len(reader.pages)}
    if suffix == ".docx":
        from docx import Document
        from docx.text.paragraph import Paragraph

        validate_office_zip(data)
        document = Document(io.BytesIO(data))
        parts = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            else:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        return bounded_text("\n".join(parts)), {}
    if suffix in {".txt", ".md", ".json", ".log"} or content_type.startswith("text/"):
        text = decode_text(data, content_type)
        if "\x00" in text:
            raise FileTaskError("unsupported_format", "文件似乎是二进制内容，不能作为文本读取。")
        return bounded_text(text), {}
    raise FileTaskError("unsupported_format", "已保存文件，但当前只能解析 HTML、TXT、Markdown、JSON、PDF、DOCX、CSV、TSV、XLSX。")


def _safe_cell(value):
    if value is None or isinstance(value, (str, int, float, bool, Decimal)):
        if isinstance(value, (float, Decimal)) and not math.isfinite(value):
            raise FileTaskError("invalid_cell", "表格单元格不能包含 NaN 或无穷大。")
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    raise FileTaskError("invalid_cell", "表格单元格只能包含文本、数字、布尔值或 null。")


def create_document(name: str, content: str) -> bytes:
    suffix = Path(name).suffix.lower()
    if len(content) > MAX_TEXT_CHARS:
        raise FileTaskError("document_too_large", "生成内容不能超过 100 万字符。")
    if suffix == ".xlsx":
        from openpyxl import Workbook

        try:
            rows = json.loads(content)
        except ValueError:
            raise FileTaskError("invalid_rows", "生成 XLSX 时 content 必须是二维 JSON 数组，第一行是列名。") from None
        if not isinstance(rows, list) or len(rows) > MAX_TABLE_ROWS + 1 or any(not isinstance(row, list) or len(row) > MAX_TABLE_COLUMNS for row in rows):
            raise FileTaskError("invalid_rows", "请提供不超过 10000 行、100 列的二维数组。")
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("结果")
        for row in rows:
            worksheet.append([_safe_cell(value) for value in row])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()
    if suffix in {".csv", ".tsv"}:
        buffer = io.StringIO(newline="")
        delimiter = "\t" if suffix == ".tsv" else ","
        writer = csv.writer(buffer, delimiter=delimiter)
        for row in csv.reader(io.StringIO(content), delimiter=delimiter):
            writer.writerow([_safe_cell(value) for value in row])
        return buffer.getvalue().encode("utf-8-sig")
    if suffix in {".txt", ".md", ".json"}:
        if suffix == ".json":
            try:
                json.loads(content)
            except ValueError:
                raise FileTaskError("invalid_json", "文件内容不是有效的 JSON。") from None
        return content.encode("utf-8")
    raise FileTaskError("unsupported_output", "可生成 TXT、Markdown、JSON、CSV、TSV 或 XLSX 文件。")


def summarize_table(table: Table, group_by: list[str], value_columns: list[str], operation: str) -> tuple[list[list], int]:
    if operation not in {"sum", "count", "mean", "min", "max"}:
        raise FileTaskError("invalid_operation", "支持 sum、count、mean、min、max。")
    if len(set(group_by + value_columns)) != len(group_by + value_columns) or any(column not in table.columns for column in group_by + value_columns):
        raise FileTaskError("invalid_columns", "分组列和数值列必须来自表头，且不能重复。")
    if operation != "count" and not value_columns:
        raise FileTaskError("missing_columns", "请指定要统计的数值列。")
    if operation == "count" and value_columns:
        raise FileTaskError("invalid_columns", "count 按分组统计行数，value_columns 应为空。")
    group_indexes = [table.columns.index(column) for column in group_by]
    value_indexes = [table.columns.index(column) for column in value_columns]
    groups: dict[tuple, tuple[int, list[Decimal]]] = {}
    for row_number, row in enumerate(table.rows, start=2):
        key = tuple(str(row[index]) if row[index] is not None else "" for index in group_indexes)
        values = []
        for index in value_indexes:
            try:
                value = Decimal(str(row[index]).strip())
                if not value.is_finite():
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                raise FileTaskError("non_numeric_cell", f"第 {row_number} 行的 {table.columns[index]} 不是有限数字（可能为空、公式或带单位）。未生成统计结果。") from None
            values.append(value)
        count, previous = groups.get(key, (0, values))
        if count:
            if operation in {"sum", "mean"}:
                values = [a + b for a, b in zip(previous, values)]
            elif operation in {"min", "max"}:
                reducer = min if operation == "min" else max
                values = [reducer(a, b) for a, b in zip(previous, values)]
        groups[key] = (count + 1, values)
    headers = group_by + (["count"] if operation == "count" else [f"{column}_{operation}" for column in value_columns])
    result = [headers]
    for key, (count, values) in groups.items():
        numbers = [count] if operation == "count" else [float(value / count if operation == "mean" else value) for value in values]
        if any(not math.isfinite(value) for value in numbers):
            raise FileTaskError("numeric_out_of_range", "统计值超出 Excel 可表示的有限数值范围。")
        result.append([*key, *numbers])
    return result, len(table.rows)
