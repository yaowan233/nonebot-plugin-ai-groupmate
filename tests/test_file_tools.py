import io
import os
import json
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
from openpyxl import Workbook, load_workbook


def workspace(tmp_path, **overrides):
    from nonebot_plugin_ai_groupmate.agent.file_workspace import FileWorkspace

    return FileWorkspace(tmp_path, **{"session_id": "group", "user_id": "user", "bot_id": "bot", "is_private": False, **overrides})


def toolset(tmp_path, **overrides):
    from nonebot_plugin_alconna import Target

    from nonebot_plugin_ai_groupmate.agent.file_tools import create_file_tools

    request_id = overrides.pop("request_id", None)
    session = SimpleNamespace(commit=AsyncMock(), add=lambda row: None)
    kwargs = {"root": tmp_path, "user_id": "user", "bot_id": "bot", "is_private": False, "bot_name": "bot", "send_target": Target("group")}
    kwargs.update(overrides)
    return {item.name: item for item in create_file_tools(session, "group", request_id, **kwargs)}


def xlsx_bytes(rows):
    book = Workbook()
    assert book.active is not None
    for row in rows:
        book.active.append(row)
    result = io.BytesIO()
    book.save(result)
    book.close()
    return result.getvalue()


@pytest.mark.parametrize("suffix", ["html", "svg", "py", "js", "css", "yaml", "toml"])
async def test_source_files_are_saved_and_read_without_execution(tmp_path, suffix):
    tools = toolset(tmp_path)
    source = '<script>throw new Error("not executed")</script>'
    created = json.loads(await tools["create_file"].ainvoke({"filename": "sample." + suffix, "content": source}))
    assert created["status"] == "succeeded"
    result = json.loads(await tools["read_file"].ainvoke({"file_id": created["data"]["file_id"], "mode": "source"}))
    assert result["data"]["text"] == source


async def test_image_read_selects_gif_frame_and_preserves_scope(tmp_path):
    import base64

    from PIL import Image

    stream = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(stream, format="GIF", save_all=True, append_images=[Image.new("RGB", (8, 8), "blue")], duration=100, loop=0)
    artifact = workspace(tmp_path).save(stream.getvalue(), "animation.gif")
    result = await toolset(tmp_path)["read_file"].ainvoke({"file_id": artifact.file_id, "mode": "image", "frame": 1})
    assert isinstance(result, list)
    assert json.loads(result[0]["text"])["data"]["frame"] == 1
    image = Image.open(io.BytesIO(base64.b64decode(result[1]["image_url"]["url"].split(",")[1])))
    pixel = image.convert("RGB").getpixel((0, 0))
    assert isinstance(pixel, tuple)
    assert pixel[2] > 240
    denied = json.loads(await toolset(tmp_path, user_id="other")["read_file"].ainvoke({"file_id": artifact.file_id, "mode": "image"}))
    assert denied["status"] == "failed"
    missing = json.loads(await toolset(tmp_path)["read_file"].ainvoke({"file_id": artifact.file_id, "mode": "image", "frame": 3}))
    assert missing["reason_code"] == "invalid_frame"


def test_workspace_isolation_persistence_and_expiry(tmp_path, monkeypatch):
    from nonebot_plugin_ai_groupmate.agent import file_workspace as module

    store = workspace(tmp_path)
    file = store.save(b"private", "../../secret.txt")
    assert file.name == "secret.txt"
    assert workspace(tmp_path).read(file.file_id) == b"private"
    for scope in ({"session_id": "other"}, {"user_id": "other"}, {"bot_id": "other"}, {"is_private": True}):
        other = workspace(tmp_path, **scope)
        assert other.list() == []
        with pytest.raises(module.FileTaskError, match="不存在"):
            other.read(file.file_id)
    with pytest.raises(module.FileTaskError, match="路径"):
        store.read("../secret.txt")
    monkeypatch.setattr(module.time, "time", lambda: file.created_at + module.FILE_TTL_SECONDS + 10)
    assert store.list() == []
    module.cleanup_workspace(tmp_path)
    assert not (store.directory / file.file_id).exists()


def test_workspace_limits_and_cleanup_preserves_unrelated_paths(tmp_path, monkeypatch):
    from nonebot_plugin_ai_groupmate.agent import file_workspace as module

    sentinel = tmp_path / "unrelated.txt"
    sentinel.write_text("keep")
    os.utime(sentinel, (0, 0))
    store = workspace(tmp_path)
    monkeypatch.setattr(module, "MAX_SCOPE_FILES", 1)
    store.save(b"data", "first.txt")
    with pytest.raises(module.FileTaskError, match="已满"):
        store.save(b"data", "second.txt")
    with pytest.raises(module.FileTaskError, match="10 MB"):
        store.save(b"x" * (module.MAX_FILE_BYTES + 1), "large.bin")
    module.cleanup_workspace(tmp_path)
    assert sentinel.read_text() == "keep"


@pytest.mark.asyncio
async def test_http_download_pins_ip_and_validates_redirects(monkeypatch):
    from nonebot_plugin_ai_groupmate.agent import web_fetch as module

    looked_up = []
    requests = []

    def resolve(host, port, **kwargs):
        looked_up.append(host)
        return [(2, 1, 6, "", ("93.184.216.34", port))]

    def respond(request):
        requests.append(request)
        if request.url.path == "/start":
            return httpx.Response(302, headers={"location": "https://cdn.example.org/files/table.csv"})
        return httpx.Response(200, content=b"a,b\n1,2", headers={"content-type": "text/csv", "content-disposition": "attachment; filename*=UTF-8''report.csv"})

    original = httpx.AsyncClient
    monkeypatch.setattr(module.socket, "getaddrinfo", resolve)
    monkeypatch.setattr(module.httpx, "AsyncClient", lambda **kwargs: original(transport=httpx.MockTransport(respond), **kwargs))
    result = await module.download_public_url("https://example.org/start")
    assert result.name == "report.csv"
    assert result.url == "https://cdn.example.org/files/table.csv"
    assert result.data == b"a,b\n1,2"
    assert looked_up == ["example.org", "cdn.example.org"]
    assert all(request.url.host == "93.184.216.34" for request in requests)
    assert requests[0].headers["host"] == "example.org"
    assert requests[1].extensions["sni_hostname"] == "cdn.example.org"


@pytest.mark.asyncio
@pytest.mark.parametrize("addresses", [["127.0.0.1"], ["::1"], ["169.254.169.254"], ["93.184.216.34", "10.0.0.1"]])
async def test_download_blocks_private_and_mixed_dns(monkeypatch, addresses):
    from nonebot_plugin_ai_groupmate.agent import web_fetch as module

    monkeypatch.setattr(module.socket, "getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", (ip, 443)) for ip in addresses])
    with pytest.raises(module.FileTaskError, match="公开网络"):
        await module.download_public_url("https://example.org/file")


@pytest.mark.asyncio
async def test_redirect_cannot_reach_internal_network(monkeypatch):
    from nonebot_plugin_ai_groupmate.agent import web_fetch as module

    calls = []

    def resolve(host, port, **kwargs):
        return [(2, 1, 6, "", ("127.0.0.1" if host == "internal.example" else "93.184.216.34", port))]

    def respond(request):
        calls.append(request)
        return httpx.Response(302, headers={"location": "http://internal.example/secret"})

    original = httpx.AsyncClient
    monkeypatch.setattr(module.socket, "getaddrinfo", resolve)
    monkeypatch.setattr(module.httpx, "AsyncClient", lambda **kwargs: original(transport=httpx.MockTransport(respond), **kwargs))
    with pytest.raises(module.FileTaskError, match="公开网络"):
        await module.download_public_url("https://example.org/file")
    assert len(calls) == 1


@pytest.mark.parametrize("url", ["file:///etc/passwd", "http://user:pass@example.org", "https://example.org:9000", "https://example.org/\nsecret"])
def test_invalid_download_urls(url):
    from nonebot_plugin_ai_groupmate.agent.web_fetch import public_url
    from nonebot_plugin_ai_groupmate.agent.file_workspace import FileTaskError

    with pytest.raises(FileTaskError):
        public_url(url)


@pytest.mark.asyncio
async def test_chunked_body_limit_without_content_length(monkeypatch):
    from nonebot_plugin_ai_groupmate.agent import web_fetch as module

    class Body(httpx.AsyncByteStream):
        async def __aiter__(self):
            yield b"a" * 100
            yield b"b" * 100

    original = httpx.AsyncClient
    monkeypatch.setattr(module, "MAX_FILE_BYTES", 150)
    monkeypatch.setattr(module.socket, "getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("93.184.216.34", 443))])
    monkeypatch.setattr(module.httpx, "AsyncClient", lambda **kwargs: original(transport=httpx.MockTransport(lambda request: httpx.Response(200, stream=Body())), **kwargs))
    with pytest.raises(module.FileTaskError, match="10 MB"):
        await module.download_public_url("https://example.org/file")


def test_html_extraction_keeps_tables_links_and_drops_scripts():
    from nonebot_plugin_ai_groupmate.agent.file_content import extract_html

    title, text, links = extract_html(b'<title>Prices</title><script>steal()</script><table><tr><th>A</th><th>B</th></tr><tr><td>2</td><td>3</td></tr></table><a href="../data.xlsx">Download</a>', "https://example.org/news/page")
    assert title == "Prices"
    assert "steal" not in text
    assert "A | B" in text
    assert "2 | 3" in text
    assert links == [{"text": "Download", "url": "https://example.org/data.xlsx"}]


@pytest.mark.asyncio
async def test_long_documents_are_contiguously_readable(tmp_path):
    tools = toolset(tmp_path)
    original = '甲\n乙"' * 1600
    file = json.loads(await tools["create_file"].ainvoke({"filename": "long.txt", "content": original}))["data"]
    pieces, offset = [], 0
    while offset is not None:
        raw = await tools["read_file"].ainvoke({"file_id": file["file_id"], "offset": offset})
        assert len(raw) < 6000
        result = json.loads(raw)["data"]
        pieces.append(result["text"])
        assert result["safety_notice"]
        offset = result["next_offset"]
    assert "".join(pieces) == original
    assert json.loads(await toolset(tmp_path)["list_files"].ainvoke({}))["data"]["files"][0]["file_id"] == file["file_id"]


def test_docx_and_text_pdf_extraction():
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject

    from nonebot_plugin_ai_groupmate.agent.file_content import extract_document

    document = Document()
    document.add_paragraph("Report")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Total"
    table.cell(0, 1).text = "42"
    output = io.BytesIO()
    document.save(output)
    assert "Total | 42" in extract_document(output.getvalue(), "report.docx")[0]

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Total 42) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    assert "Total 42" in extract_document(output.getvalue(), "report.pdf")[0]


def test_office_expansion_limit(monkeypatch):
    from nonebot_plugin_ai_groupmate.agent import file_content as module

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", "a" * 1000)
    monkeypatch.setattr(module, "MAX_UNPACKED_BYTES", 100)
    with pytest.raises(module.FileTaskError, match="解压"):
        module.validate_office_zip(output.getvalue())


@pytest.mark.asyncio
async def test_table_summary_checks_all_rows_and_refuses_formulas(tmp_path):
    tools = toolset(tmp_path)
    store = workspace(tmp_path)
    original = store.save(xlsx_bytes([["类别", "金额"], ["A", 0.1], ["B", -3], ["A", 0.2]]), "data.xlsx")
    result = json.loads(await tools["summarize_table"].ainvoke({"file_id": original.file_id, "group_by": ["类别"], "value_columns": ["金额"]}))
    assert result["status"] == "succeeded"
    assert result["data"]["preview"] == [["A", 0.3], ["B", -3.0]]
    assert result["data"]["input_rows"] == 3
    generated = load_workbook(io.BytesIO(store.read(result["data"]["file_id"])), data_only=False)
    assert generated.active is not None
    assert generated.active["B3"].value == -3
    generated.close()
    invalid = store.save(xlsx_bytes([["类别", "金额"], ["A", "=1+2"]]), "formula.xlsx")
    failure = json.loads(await tools["summarize_table"].ainvoke({"file_id": invalid.file_id, "group_by": ["类别"], "value_columns": ["金额"]}))
    assert failure["reason_code"] == "non_numeric_cell"


@pytest.mark.asyncio
async def test_creation_never_executes_spreadsheet_formulas(tmp_path):
    tools = toolset(tmp_path)
    data = json.loads(await tools["create_file"].ainvoke({"filename": "report.xlsx", "content": json.dumps([["value"], ['=WEBSERVICE("https://example.org")']])}))["data"]
    book = load_workbook(io.BytesIO(workspace(tmp_path).read(data["file_id"])))
    assert book.active is not None
    assert book.active["A2"].data_type == "s"
    book.close()


@pytest.mark.asyncio
async def test_send_unknown_delivery_is_not_retried(tmp_path, monkeypatch):
    from nonebot_plugin_alconna import UniMessage

    tools = toolset(tmp_path)
    created = json.loads(await tools["create_file"].ainvoke({"filename": "result.txt", "content": "result"}))["data"]
    sender = AsyncMock(side_effect=RuntimeError("adapter unavailable"))
    monkeypatch.setattr(UniMessage, "send", sender)
    result = json.loads(await tools["send_file"].ainvoke({"file_id": created["file_id"]}))
    assert result["status"] == "failed"
    assert result["delivery_state"] == "unknown"
    second = json.loads(await tools["send_file"].ainvoke({"file_id": created["file_id"]}))
    assert second["reason_code"] == "duplicate_file"
    assert sender.await_count == 1


@pytest.mark.asyncio
async def test_expired_request_cannot_download_or_send(tmp_path, monkeypatch):
    from nonebot_plugin_alconna import UniMessage

    from nonebot_plugin_ai_groupmate.agent import file_tools

    download = AsyncMock()
    sender = AsyncMock()
    monkeypatch.setattr(file_tools, "download_public_url", download)
    monkeypatch.setattr(UniMessage, "send", sender)
    tools = toolset(tmp_path, request_id="expired-request")
    for name, args in [("fetch_url", {"url": "https://example.org"}), ("send_file", {"file_id": "0" * 32})]:
        result = json.loads(await tools[name].ainvoke(args))
        assert result["reason_code"] == "request_expired"
    download.assert_not_called()
    sender.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize("enabled", [True, False])
@pytest.mark.parametrize("private", [True, False])
async def test_file_tools_are_wired_into_group_and_private_agents(monkeypatch, enabled, private):
    from nonebot_plugin_ai_groupmate import agent

    captured = []
    monkeypatch.setattr(agent.plugin_config, "agent_file_tools_enabled", enabled)
    for name in ("get_user_relation_context", "get_group_context", "get_recent_relations_context"):
        monkeypatch.setattr(agent, name, AsyncMock(return_value=""))
    monkeypatch.setattr(agent, "build_registered_agent_extensions", AsyncMock(return_value=([], [], [])))
    monkeypatch.setattr(agent, "get_chat_model", lambda: object())

    def build(model, tools, system_prompt, **kwargs):
        captured.extend(item.name for item in kwargs["base_tools"])
        return object()

    monkeypatch.setattr(agent, "build_chat_graph", build)
    await agent.create_chat_graph(SimpleNamespace(commit=AsyncMock()), "file-integration-group", None, "file-user", "User", history=[], is_private=private)
    assert ("fetch_url" in captured) is enabled
    assert ("send_file" in captured) is enabled


@pytest.mark.asyncio
async def test_full_agent_fetch_download_summarize_and_send(tmp_path, monkeypatch):
    """Exercise real graph + HTTP + file parsers, stubbing only transport/model/channel."""
    from nonebot_plugin_alconna import UniMessage
    from langchain_core.messages import AIMessage, ToolMessage, HumanMessage

    from nonebot_plugin_ai_groupmate.agent import web_fetch
    from nonebot_plugin_ai_groupmate.agent.graph import build_chat_graph
    from nonebot_plugin_ai_groupmate.agent.common_tools import finish

    spreadsheet = xlsx_bytes([["Category", "Amount"], ["A", 10], ["B", 5], ["A", 20]])

    def respond(request):
        if request.url.path == "/page":
            return httpx.Response(200, headers={"content-type": "text/html"}, content=b'<h1>Data</h1><a href="/data.xlsx">Excel</a>')
        assert request.url.path == "/data.xlsx"
        return httpx.Response(200, content=spreadsheet, headers={"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})

    original = httpx.AsyncClient
    monkeypatch.setattr(web_fetch.socket, "getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("93.184.216.34", 443))])
    monkeypatch.setattr(web_fetch.httpx, "AsyncClient", lambda **kwargs: original(transport=httpx.MockTransport(respond), **kwargs))
    sent = []

    async def send(message, **kwargs):
        sent.append((message, kwargs))
        return SimpleNamespace(msg_ids=[{"message_id": "42"}])

    monkeypatch.setattr(UniMessage, "send", send)

    class Model:
        def __init__(self):
            self.step = 0

        def bind_tools(self, tools):
            return self

        async def ainvoke(self, messages, **kwargs):
            results = [json.loads(message.content) for message in messages if isinstance(message, ToolMessage) and isinstance(message.content, str) and message.content]
            if results:
                assert results[-1]["status"] == "succeeded"
            if self.step == 0:
                name, args = "fetch_url", {"url": "https://example.org/page"}
            elif self.step == 1:
                name, args = "download_file", {"url": results[-1]["data"]["links"][0]["url"]}
            elif self.step == 2:
                self.input_id = results[-1]["data"]["file_id"]
                name, args = "read_file", {"file_id": self.input_id}
            elif self.step == 3:
                assert results[-1]["data"]["columns"] == ["Category", "Amount"]
                name, args = "summarize_table", {"file_id": self.input_id, "group_by": ["Category"], "value_columns": ["Amount"]}
            elif self.step == 4:
                self.output_id = results[-1]["data"]["file_id"]
                name, args = "read_file", {"file_id": self.output_id}
            elif self.step == 5:
                assert "A,30" in results[-1]["data"]["text"]
                name, args = "send_file", {"file_id": self.output_id}
            else:
                name, args = "finish", {}
            self.step += 1
            return AIMessage(content="", tool_calls=[{"id": f"call-{self.step}", "name": name, "args": args}])

    model = Model()
    tools = [*toolset(tmp_path).values(), finish]
    graph = build_chat_graph(model, tools, "Read the web page, download Excel, summarize and deliver.")
    result = await graph.ainvoke({"messages": [HumanMessage(content="Summarize the Excel on this web page and send me the file.")], "session_id": "group", "request_id": None})
    assert result["called_finish"] == 1
    assert model.step == 7
    assert len(sent) == 1
    assert sent[0][1]["fallback"] is False
    assert sent[0][1]["target"].id == "group"
    segment = sent[0][0][0]
    book = load_workbook(io.BytesIO(segment.path.read_bytes()))
    assert book.active is not None
    assert list(book.active.values) == [("Category", "Amount_sum"), ("A", 30), ("B", 5)]
    book.close()


@pytest.mark.asyncio
async def test_connected_group_bots_can_deliver_existing_files_without_crossing_users(tmp_path, monkeypatch):
    from nonebot_plugin_ai_groupmate.agent import file_workspace as module
    import nonebot

    monkeypatch.setattr(nonebot, "get_bots", lambda: {"bot": object(), "peer": object()})
    old = workspace(tmp_path).save(b"picture", "desert_boss.png")
    peer = workspace(tmp_path, bot_id="peer")
    assert peer.read(old.file_id) == b"picture"
    assert old.file_id in [item.file_id for item in peer.list()]
    for scope in ({"user_id": "other"}, {"session_id": "other"}, {"bot_id": "unconnected"}, {"is_private": True}):
        with pytest.raises(module.FileTaskError):
            workspace(tmp_path, **{"bot_id": "peer", **scope}).read(old.file_id)
    tools = toolset(tmp_path, bot_id="peer")
    from nonebot_plugin_alconna import UniMessage

    sender = AsyncMock(return_value=SimpleNamespace(msg_ids=[{"message_id": "42"}]))
    monkeypatch.setattr(UniMessage, "send", sender)
    result = json.loads(await tools["send_file"].ainvoke({"file_id": old.file_id}))
    assert result["status"] == "succeeded"
    sender.assert_awaited_once()
